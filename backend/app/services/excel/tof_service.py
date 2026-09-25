# backend/app/services/excel/tof_service.py
"""
TOF (Training Order Form) Excel Service - Stub/Template
Tác giả: Nguyễn Mạnh Hùng & Co-pilot AI
Ghi chú: Khung dịch vụ bóc tách định dạng TOF phục vụ các lớp Training / Demo.
"""
import os
import re
import gc
import logging
from typing import Dict, Any
import openpyxl

logger = logging.getLogger(__name__)


class TOFExcelService:
    """Service xử lý file TOF (Training Order Form)."""

    @classmethod
    def is_tof_file(cls, file_path: str) -> bool:
        """
        Nhận diện chính xác file TOF:
        - Kiểm tra tên file (fname).
        - Quét 1-5 hàng đầu, cột A-Z của Sheet 0 để tìm 'PYTHAVERSE TRAINING ORDER FORM (TOF)'
          (bảo đảm bắt dính ô merge C1:R3 dù có version hay ký tự thừa).
        """
        fname = os.path.basename(file_path).lower()
        if "tof" in fname or "training order form" in fname:
            return True

        wb = None
        try:
            wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
            sheet_names = wb.sheetnames
            if not sheet_names:
                return False

            first_sheet = wb[sheet_names[0]]
            text_parts = []
            for row in first_sheet.iter_rows(min_row=1, max_row=5, min_col=1, max_col=26, values_only=True):
                for cell in row:
                    if cell is not None:
                        val = str(cell).strip()
                        if val:
                            text_parts.append(val)

            combined = re.sub(r"\s+", " ", " ".join(text_parts)).upper()
            if "TRAINING ORDER FORM" in combined or "(TOF)" in combined:
                return True

            # Kiểm tra tên tab
            for sname in sheet_names:
                if "TRAINING ORDER FORM" in sname.upper() or "(TOF)" in sname.upper():
                    return True

        except Exception as e:
            logger.warning(f"⚠️ Lỗi đọc header TOF: {e}")
        finally:
            if wb:
                wb.close()
                del wb
                gc.collect()

        return False

    @classmethod
    def parse_tof_file(cls, file_path: str) -> Dict[str, Any]:
        """Tạm thời trả về khung dữ liệu chuẩn của TOF (Anh sẽ hướng dẫn chi tiết sau)."""
        logger.info(f"📄 [TOF SERVICE] Đang bóc tách file TOF: {file_path}")
        return {
            "file_type": "tof",
            "filename": os.path.basename(file_path),
            "status": "ready_for_specification",
            "message": "Đang chờ cập nhật quy chuẩn bóc tách TOF từ Kiến trúc sư Nguyễn Mạnh Hùng."
        }

    @classmethod
    def parse_tof_summary(cls, file_path: str) -> Dict[str, Any]:
        """Bóc tách summary của TOF phục vụ AI Intake."""
        data = cls.parse_tof_file(file_path)
        return {
            "school_name": data.get("school_name", ""),
            "courses_detected": data.get("courses_detected", [])
        }


tof_service = TOFExcelService()