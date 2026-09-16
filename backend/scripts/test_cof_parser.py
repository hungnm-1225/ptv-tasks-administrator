# backend/test_cof_parser.py
"""
=====================================================================
🧪 SCRIPT TEST LOCAL BÓC TÁCH FILE COF CHUẨN XÁC TỪNG TỌA ĐỘ
=====================================================================
Tự động tìm file COF trong thư mục /backend/ và bóc tách chuẩn xác:
- Tab 1: Curriculum Order Form (COF)
  * Trường & Địa chỉ: Cột C hàng 6 & 7
  * Tên môn: CỘT G (Cột 7)
  * Course ID: CỘT H (Cột 8)
  * Ngày bắt đầu: CỘT I (Cột 9 - merge I, J, K)
  * Ngày kết thúc: CỘT L (Cột 12 - merge L, M, N)
  * Ngày yêu cầu: CỘT O (Cột 15 - merge O, P)
  * Số lượng bản quyền: CỘT Q (Cột 17) -> ĐÚNG 100%!
  * LỌC CỐT TỬ: Chỉ lấy môn có ít nhất 1 trong 3 (Start Date, End Date, Quantity).
                Môn nào thiếu cả 3 -> LOẠI BỎ NGAY (môn mẫu)!
- Tab 2: Student Information (Học sinh)
- Tab 3: Teacher Information (Giáo viên)
=====================================================================
"""

import os
import sys
import re
import json
import logging
from datetime import datetime, date, timedelta
from typing import Dict, Any, List, Optional
from pathlib import Path
import openpyxl

logging.basicConfig(level=logging.INFO, format="%(asctime)s [%(levelname)s] %(message)s")
logger = logging.getLogger("COF_FAST_TEST")


def format_date_str(val: Any) -> str:
    """Chuẩn hóa ngày tháng về YYYY-MM-DD an toàn, xử lý cả Excel Serial Number."""
    if not val or str(val).strip() in ["", "None", "null", "---"]:
        return ""
    
    if isinstance(val, (datetime, date)):
        return val.strftime("%Y-%m-%d")

    # Xử lý số ngày Excel (Serial Date Number)
    if isinstance(val, (int, float)):
        try:
            val_num = float(val)
            if val_num > 1000:
                d = datetime(1899, 12, 30) + timedelta(days=val_num)
                return d.strftime("%Y-%m-%d")
        except Exception:
            pass

    val_str = str(val).strip().split(" ")[0].split("T")[0]
    
    # Định dạng DD-MM-YYYY hoặc DD/MM/YYYY
    d_match = re.match(r"^(\d{1,2})[-/\.](\d{1,2})[-/\.](\d{4})$", val_str)
    if d_match:
        d, m, y = d_match.groups()
        return f"{y}-{int(m):02d}-{int(d):02d}"

    # Định dạng YYYY-MM-DD hoặc YYYY/MM/DD
    y_match = re.match(r"^(\d{4})[-/\.](\d{1,2})[-/\.](\d{1,2})$", val_str)
    if y_match:
        y, m, d = y_match.groups()
        return f"{y}-{int(m):02d}-{int(d):02d}"

    if "dd-mm-yyyy" in val_str.lower():
        return ""

    return val_str


def parse_clean_number(val: Any) -> int:
    """Bóc tách số nguyên sạch từ ô tính, nếu không có trả về 0."""
    if val is None:
        return 0
    val_str = str(val).strip()
    match = re.search(r"\d+", val_str)
    if match:
        return int(match.group(0))
    return 0


def detect_category(course_name: str) -> str:
    """Nhận diện Category chuẩn xác từ tên môn học."""
    c_upper = course_name.upper()
    if "ASP" in c_upper:
        return "ASP"
    if "IR" in c_upper or "PYTHON" in c_upper or "TERRA" in c_upper:
        return "IR"
    if "SWRP" in c_upper:
        return "SWRP"
    if "LEANBOT" in c_upper or "C++" in c_upper:
        return "Other"
    return "SWRP"


def find_target_cof_file() -> str:
    """Tự động quét tìm file COF trong thư mục backend hoặc thư mục hiện tại."""
    # 1. Nếu truyền qua tham số dòng lệnh
    if len(sys.argv) > 1 and os.path.exists(sys.argv[1]):
        return sys.argv[1]

    # 2. Quét trong backend/ hoặc root
    search_dirs = [
        Path(__file__).resolve().parent,
        Path(__file__).resolve().parent / "data" / "cof_input",
        Path.cwd(),
        Path.cwd() / "backend"
    ]

    for d in search_dirs:
        if d.exists():
            for f in os.listdir(d):
                if f.endswith(".xlsx") and not f.startswith("~"):
                    f_lower = f.lower()
                    if "cof" in f_lower or "curriculum" in f_lower or "maasin" in f_lower or "result" in f_lower:
                        return str(d / f)

    return ""


def test_parse_cof(file_path: str):
    logger.info("=" * 70)
    logger.info(f"📂 ĐANG BÓC TÁCH FILE COF: {os.path.basename(file_path)}")
    logger.info("=" * 70)

    wb = openpyxl.load_workbook(file_path, data_only=True)
    try:
        sheet_names = wb.sheetnames

        # 1. TAB 1: CURRICULUM ORDER FORM
        tab1_name = next((s for s in sheet_names if "curriculum" in s.lower() or "cof" in s.lower()), sheet_names[0])
        ws1 = wb[tab1_name]

        # Bóc tách Header trường học
        school_name = str(ws1.cell(row=6, column=3).value or "").strip()
        school_address = str(ws1.cell(row=7, column=3).value or "").strip()
        tel = str(ws1.cell(row=6, column=9).value or ws1.cell(row=6, column=8).value or "").strip()
        email = str(ws1.cell(row=7, column=9).value or ws1.cell(row=7, column=8).value or "").strip()

        logger.info(f"🏫 Tên trường bóc tách: '{school_name}'")
        logger.info(f"📍 Địa chỉ trường: '{school_address}'")

        # Bóc tách bảng môn học
        ordered_courses = []
        sample_ignored_count = 0

        # Bảng môn học bắt đầu từ hàng 28
        for r in range(28, ws1.max_row + 1):
            course_link = str(ws1.cell(row=r, column=3).value or "").strip()   # Cột C (Course Link)
            course_name = str(ws1.cell(row=r, column=7).value or "").strip()   # Cột G (Course Name thật)
            raw_course_id = ws1.cell(row=r, column=8).value                   # Cột H (Course ID thật)
            
            start_date_raw = ws1.cell(row=r, column=9).value                  # Cột I (Start Date - merge I-K)
            end_date_raw = ws1.cell(row=r, column=12).value                   # Cột L (End Date - merge L-N)
            date_req_raw = ws1.cell(row=r, column=15).value                   # Cột O (Date Request - merge O-P)
            qty_raw = ws1.cell(row=r, column=17).value                        # Cột Q (Số lượng bản quyền thật!)

            # Bỏ qua dòng tiêu đề nhóm hoặc dòng trống
            if not course_name and not raw_course_id:
                continue
            if "course name" in course_name.lower() or "courses" in course_name.lower() and not raw_course_id:
                continue

            course_id = parse_clean_number(raw_course_id)
            start_date = format_date_str(start_date_raw)
            end_date = format_date_str(end_date_raw)
            date_request = format_date_str(date_req_raw)
            quantity = parse_clean_number(qty_raw)

            # 🎯 ĐIỀU KIỆN LỌC CỐT TỬ CỦA ANH:
            # Chỉ nhận môn có ÍT NHẤT 1 trong 3 thông tin: Start Date, End Date, Quantity.
            has_start = bool(start_date)
            has_end = bool(end_date)
            has_qty = quantity > 0

            if not (has_start or has_end or has_qty):
                sample_ignored_count += 1
                continue

            category = detect_category(course_name)

            ordered_courses.append({
                "row_index": r,
                "course_id": str(course_id) if course_id else "",
                "course_name": course_name,
                "course_link": course_link,
                "category": category,
                "licenses": quantity,                 # ĐỌC TỪ CỘT Q CHUẨN XÁC!
                "start_date": start_date or None,     # CỘT I
                "end_date": end_date or None,         # CỘT L
                "date_request": date_request or None  # CỘT O
            })

        logger.info(f"✨ BÓC TÁCH THÀNH CÔNG {len(ordered_courses)} MÔN THỰC ĐẶT (Đã loại bỏ {sample_ignored_count} môn mẫu rác).")

        # 2. TAB 2: STUDENT INFORMATION
        tab2_name = next((s for s in sheet_names if "student" in s.lower()), None)
        students_all = []
        students_to_create = []

        if tab2_name:
            ws2 = wb[tab2_name]
            for r in range(7, ws2.max_row + 1):
                full_name = str(ws2.cell(row=r, column=2).value or "").strip()
                first_name = str(ws2.cell(row=r, column=3).value or "").strip()
                last_name = str(ws2.cell(row=r, column=4).value or "").strip()
                email = str(ws2.cell(row=r, column=6).value or "").strip().lower()
                dob_raw = ws2.cell(row=r, column=7).value
                account_exist = str(ws2.cell(row=r, column=8).value or "").strip().lower()
                grade = str(ws2.cell(row=r, column=9).value or "").strip()
                class_group = str(ws2.cell(row=r, column=11).value or "").strip()
                username = str(ws2.cell(row=r, column=12).value or "").strip()

                if not first_name and not full_name:
                    continue
                if "total" in full_name.lower():
                    continue

                student_rec = {
                    "row_index": r,
                    "full_name": full_name or f"{first_name} {last_name}".strip(),
                    "first_name": first_name or (full_name.split()[0] if full_name else "Student"),
                    "last_name": last_name or (" ".join(full_name.split()[1:]) if len(full_name.split()) > 1 else "Auto"),
                    "email": email,
                    "dob": format_date_str(dob_raw) or "2016-01-01",
                    "grade": grade,
                    "class_group": class_group or "Default Class",
                    "username": username,
                    "role": "student"
                }
                students_all.append(student_rec)
                if not username and account_exist != "yes":
                    students_to_create.append(student_rec)

        # 3. TAB 3: TEACHER INFORMATION
        tab3_name = next((s for s in sheet_names if "teacher" in s.lower()), None)
        teachers_all = []
        teachers_to_create = []

        if tab3_name:
            ws3 = wb[tab3_name]
            for r in range(7, ws3.max_row + 1):
                class_group = str(ws3.cell(row=r, column=3).value or "").strip()
                full_name = str(ws3.cell(row=r, column=5).value or "").strip()
                first_name = str(ws3.cell(row=r, column=6).value or "").strip()
                last_name = str(ws3.cell(row=r, column=7).value or "").strip()
                email = str(ws3.cell(row=r, column=8).value or "").strip().lower()
                dob_raw = ws3.cell(row=r, column=9).value
                account_exist = str(ws3.cell(row=r, column=10).value or "").strip().lower()
                course_assign = str(ws3.cell(row=r, column=11).value or "").strip()
                username = str(ws3.cell(row=r, column=12).value or "").strip()

                if not first_name and not full_name:
                    continue
                if "total" in full_name.lower():
                    continue

                teacher_rec = {
                    "row_index": r,
                    "full_name": full_name or f"{first_name} {last_name}".strip(),
                    "first_name": first_name or (full_name.split()[0] if full_name else "Teacher"),
                    "last_name": last_name or "Auto",
                    "email": email,
                    "dob": format_date_str(dob_raw) or "1990-01-01",
                    "class_group": class_group or "Default Class",
                    "course_assign": course_assign,
                    "username": username,
                    "role": "teacher"
                }
                teachers_all.append(teacher_rec)
                if not username and account_exist != "yes":
                    teachers_to_create.append(teacher_rec)

        # IN BẢNG TỔNG KẾT BẰNG MẮT
        print("\n" + "=" * 80)
        print("🏆 BẢNG TỔNG HỢP DỮ LIỆU ĐÃ BÓC TÁCH TỪ FILE COF:")
        print("=" * 80)
        print(f"🏫 Trường học : {school_name}")
        print(f"📍 Địa chỉ    : {school_address}")
        print(f"📚 Môn thực đặt: {len(ordered_courses)} môn (Đã lọc bỏ {sample_ignored_count} môn mẫu rác)")
        print(f"🎓 Học sinh   : {len(students_all)} tổng (Cần tạo mới: {len(students_to_create)})")
        print(f"👨‍🏫 Giáo viên  : {len(teachers_all)} tổng (Cần tạo mới: {len(teachers_to_create)})")
        
        print("\n📋 DANH SÁCH KHÓA HỌC THỰC TẾ (ORDERED COURSES):")
        print(f"{'STT':<4} | {'CAT':<6} | {'ID (Cột H)':<10} | {'TÊN KHÓA HỌC (Cột G)':<40} | {'SỐ LƯỢNG (Cột Q)':<16} | {'TỪ NGÀY':<10} | {'ĐẾN NGÀY':<10}")
        print("-" * 110)
        for idx, c in enumerate(ordered_courses, 1):
            print(f"{idx:<4} | {c['category']:<6} | {c['course_id']:<10} | {c['course_name'][:38]:<40} | {c['licenses']:<16} | {str(c['start_date']):<10} | {str(c['end_date']):<10}")
        print("=" * 80 + "\n")

    finally:
        wb.close()


if __name__ == "__main__":
    target_file = find_target_cof_file()
    if not target_file:
        logger.error("❌ Không tìm thấy file Excel COF nào trong thư mục backend!")
        sys.exit(1)

    test_parse_cof(target_file)